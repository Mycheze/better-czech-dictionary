#!/usr/bin/env python3
"""
Regenerate data/youtube_channels.tsv from the Refold Czech resource sheet.

The sheet's "YouTube Channels" tab lists ~175 Czech channels with genre and
difficulty metadata. This script pulls it, buckets the channels by genre, and
picks N of them round-robin across buckets so the corpus stays topically wide:
without balancing, vlogs and gaming would crowd out news, food, tech and
history, and the harvested vocabulary would skew with it.

Requires: openpyxl (pip install openpyxl)

Usage:
    python3 tools/build_channel_list.py                  # 100 channels
    python3 tools/build_channel_list.py --count 150
    python3 tools/build_channel_list.py --out other.tsv
"""

import argparse
import csv
import re
import sys
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = PROJECT_ROOT / "data" / "youtube_channels.tsv"

SHEET_ID = "1ErJLW9BGRYD3X2PueZvbgnvfXe4rl9cIREN6FW-7Fl0"
SHEET_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
TAB = "YouTube Channels"

# Genre labels from the sheet, grouped into broad buckets for round-robin.
BUCKETS = [
    ("learner",   ["comprehensible input", "teaching", "czech language", "language learning",
                   "easy languages", "czech input", "street interviews", "czech teaching",
                   "language"]),
    ("gaming",    ["video games", "gameplay", "minecraft"]),
    ("vlog",      ["vlog", "daily life", "life in", "random videos"]),
    ("travel",    ["travel", "outdoors", "camping", "prague", "czech republic"]),
    ("tech",      ["technology", "tech", "computers", "gadgets", "space", "science"]),
    ("food",      ["food", "cooking"]),
    ("news",      ["news", "documentaries", "history", "architecture", "infotainment"]),
    ("culture",   ["pop culture", "movies", "television", "books", "audiobooks",
                   "entertainment", "comedy", "music"]),
    ("reaction",  ["reactions", "comentary", "commentary", "youtube", "livestream", "podcast"]),
    ("lifestyle", ["gym", "fitness", "self improvement", "fashion", "makeup", "real estate",
                   "reviews", "sports", "sport", "skateboarding", "trains", "public transit",
                   "public transport"]),
]

# The sheet is formula-driven; the exported xlsx keeps the computed value as the
# IFERROR fallback string, which is what we actually want to read.
CACHED_VALUE = re.compile(r'^=IFERROR\(.*,"(.*)"\)$', re.S)


def cell_value(v):
    if v is None:
        return ""
    s = str(v)
    m = CACHED_VALUE.match(s)
    return m.group(1).replace('""', '"') if m else s


def bucket_of(genre):
    for part in (p.strip() for p in genre.lower().split(",")):
        if not part:
            continue
        for name, keys in BUCKETS:
            if any(k == part or k in part for k in keys):
                return name
    return "other"


def normalize_url(url):
    """Return a yt-dlp-friendly uploads URL, or None if this isn't a channel.

    A few sheet rows point at a single video or a playlist rather than a
    channel. Appending '/videos' to those produces nonsense like
    'youtube.com/watch/videos', so playlists are passed through untouched and
    single-video links are rejected (resolve them by hand if you want them).
    """
    url = url.strip().rstrip("/")
    if "/playlist?" in url:
        return url
    if "/watch" in url:
        return None
    url = re.sub(r"\?.*$", "", url)
    for suffix in ("/videos", "/featured", "/streams", "/shorts", "/about", "/playlists"):
        if url.endswith(suffix):
            url = url[: -len(suffix)]
    if not re.search(r"youtube\.com/(@|c/|channel/|user/)", url):
        return None
    return url + "/videos"


def main():
    ap = argparse.ArgumentParser(description="Rebuild the YouTube channel list")
    ap.add_argument("--count", type=int, default=100, help="How many channels to keep")
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("ERROR: pip install openpyxl")
        sys.exit(1)

    print(f"Fetching sheet {SHEET_ID} ...")
    tmp = Path(args.out).parent / ".sheet_tmp.xlsx"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    urllib.request.urlretrieve(SHEET_URL, tmp)

    wb = openpyxl.load_workbook(tmp)
    if TAB not in wb.sheetnames:
        print(f"ERROR: tab {TAB!r} not found. Tabs: {wb.sheetnames}")
        sys.exit(1)
    ws = wb[TAB]

    rows = []
    for row in ws.iter_rows(values_only=True):
        vals = [cell_value(c) for c in row]
        if any(v.strip() for v in vals):
            rows.append(vals)
    tmp.unlink(missing_ok=True)

    try:
        hdr = next(i for i, r in enumerate(rows) if r and r[0] == "Title")
    except StopIteration:
        print("ERROR: header row not found")
        sys.exit(1)

    channels, skipped = [], 0
    for r in rows[hdr + 1:]:
        if len(r) < 4 or not r[1].startswith("http"):
            continue
        url = normalize_url(r[1])
        if not url:
            skipped += 1
            continue
        channels.append({"title": r[0].strip(), "url": url, "bucket": bucket_of(r[2]),
                         "genre": r[2].strip(), "difficulty": r[3].strip()})

    # De-duplicate on URL, keeping sheet order.
    seen, uniq = set(), []
    for c in channels:
        if c["url"] not in seen:
            seen.add(c["url"])
            uniq.append(c)

    print(f"Parsed {len(uniq)} channels ({skipped} non-channel URLs skipped)")

    by_bucket = defaultdict(list)
    for c in uniq:
        by_bucket[c["bucket"]].append(c)

    selected, i = [], 0
    order = sorted(by_bucket, key=lambda b: -len(by_bucket[b]))
    while len(selected) < args.count:
        added = False
        for b in order:
            if i < len(by_bucket[b]):
                selected.append(by_bucket[b][i])
                added = True
                if len(selected) >= args.count:
                    break
        if not added:
            break
        i += 1

    with open(args.out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, delimiter="\t",
                           fieldnames=["title", "url", "bucket", "genre", "difficulty"])
        w.writeheader()
        w.writerows(selected)

    print(f"Selected {len(selected)} channels -> {args.out}")
    print("Bucket spread:", Counter(c["bucket"] for c in selected).most_common())
    print("\nNote: channels that 404 (renamed or deleted handles) are reported by")
    print("download_yt_subs.py --status; replace them by hand or raise --count.")


if __name__ == "__main__":
    main()
